#!/usr/bin/env python3
"""
deck-builder — build an editable .pptx (sales or training deck) from a JSON
outline. Output is a working file for the owner to edit; it is never
published or sent by this tool.

Outline format (JSON):
{
  "title": "Deck title", "subtitle": "optional", "kind": "sales" | "training",
  "author": "Name", "footer": "Confidential — draft",
  "slides": [
    {"title": "Slide title", "bullets": ["point", ["sub-point", "sub-point"]], "notes": "speaker notes"},
    {"title": "Two columns", "left": ["a","b"], "right": ["c","d"], "notes": "..."},
    {"section": "Section divider title"}
  ]
}

Usage:
  python3 agents/skills/deck-builder/build_deck.py outline.json [--out path.pptx] [--engagement eng-xxx --actor agent:production-agent]
If --engagement is given, the deck is registered in the ledger as a draft deliverable.
"""
import argparse
import datetime as _dt
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
sys.path.insert(0, os.path.join(REPO, "agents", "skills", "business-records-ledger"))
import ledger  # noqa: E402

THEME = {"sales": RGBColor(0x1F, 0x3B, 0x57), "training": RGBColor(0x2E, 0x6B, 0x3A)}


def _fill_bullets(tf, items, level=0):
    first = True
    for it in items:
        if isinstance(it, list):
            _fill_bullets(tf, it, level + 1)
            continue
        p = tf.paragraphs[0] if first and level == 0 and not tf.paragraphs[0].text else tf.add_paragraph()
        first = False
        p.text = str(it)
        p.level = level
        p.font.size = Pt(20 - 3 * level)


def build(outline, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    accent = THEME.get(outline.get("kind", "sales"), THEME["sales"])
    footer = outline.get("footer", "Draft — internal working file")

    def add_footer(slide, n):
        tb = slide.shapes.add_textbox(Inches(0.4), Inches(7.0), Inches(12.5), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = f"{footer}   ·   {n}"
        p.font.size = Pt(10); p.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    def bar(slide):
        shp = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.18))
        shp.fill.solid(); shp.fill.fore_color.rgb = accent; shp.line.fill.background()

    # title slide
    s = prs.slides.add_slide(prs.slide_layouts[0])
    bar(s)
    s.shapes.title.text = outline.get("title", "Untitled deck")
    if len(s.placeholders) > 1:
        s.placeholders[1].text = outline.get("subtitle", "") + (f"\n{outline.get('author')}" if outline.get("author") else "") + f"\n{_dt.date.today().isoformat()}"
    add_footer(s, 1)

    n = 1
    for sl in outline.get("slides", []):
        n += 1
        if "section" in sl:
            s = prs.slides.add_slide(prs.slide_layouts[2])
            bar(s); s.shapes.title.text = sl["section"]
            if len(s.placeholders) > 1: s.placeholders[1].text = sl.get("subtitle", "")
        elif "left" in sl or "right" in sl:
            s = prs.slides.add_slide(prs.slide_layouts[3])
            bar(s); s.shapes.title.text = sl.get("title", "")
            _fill_bullets(s.placeholders[1].text_frame, sl.get("left", []))
            _fill_bullets(s.placeholders[2].text_frame, sl.get("right", []))
        else:
            s = prs.slides.add_slide(prs.slide_layouts[1])
            bar(s); s.shapes.title.text = sl.get("title", "")
            _fill_bullets(s.placeholders[1].text_frame, sl.get("bullets", []))
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]
        add_footer(s, n)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    return n


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("outline")
    ap.add_argument("--out")
    ap.add_argument("--engagement")
    ap.add_argument("--actor", default="agent:production-agent")
    ap.add_argument("--ledger", default=ledger.DEFAULT_LEDGER)
    a = ap.parse_args(argv)
    outline = json.load(open(a.outline, encoding="utf-8"))
    kind = outline.get("kind", "sales")
    out = a.out or os.path.join(REPO, "agents", "workspace", "40_Sales_Training", f"{kind}-deck-{_dt.date.today().isoformat()}.pptx")
    n = build(outline, out)
    print(f"deck      : {os.path.relpath(out, REPO)}  ({n} slides, kind={kind})")
    if a.engagement:
        rec = ledger.append(a.ledger, "deliverable", {"path": os.path.relpath(out, REPO), "sha256": ledger.sha256_file(out),
                                                       "status": "draft", "kind": "pptx", "bytes": os.path.getsize(out)}, a.actor, a.engagement)
        print(f"ledger rec: {rec['id']} (draft)")
    print("NOTE: editable working file. Publication / customer use requires the owner's publication gate.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
